import io
import pytest
import openpyxl
import docx
from pathlib import Path
from sqlalchemy.orm import Session

from app.models.database import Document, init_db, get_engine
from app.services.archive_service import inspect_document_content
from app.services.agent_service import AgenticChatService
from app.services.document_service import save_uploaded_file
import app.mcp_server as mcp_mod


@pytest.fixture
def db_session(tmp_path):
    db_file = tmp_path / "test_vault_excel.db"
    engine = get_engine(f"sqlite:///{db_file}")
    init_db(engine)
    with Session(engine) as session:
        yield session


def test_inspect_document_content_xlsx():
    # 1. Crea file Excel con 2 fogli
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Fatturato 2026"
    ws1.append(["Mese", "Fatturato", "Costi", "Margine"])
    ws1.append(["Gennaio", 15000, 7500, "=B2-C2"])
    ws1.append(["Febbraio", 22000, 9000, "=B3-C3"])
    ws1.append(["Marzo", 18500, 6800, "=B4-C4"])

    ws2 = wb.create_sheet(title="Clienti")
    ws2.append(["ID", "Cliente", "P.IVA"])
    ws2.append([1, "Acme Corp", "IT12345678901"])
    ws2.append([2, "Globex Srl", "IT98765432109"])

    buf = io.BytesIO()
    wb.save(buf)
    file_bytes = buf.getvalue()

    # Ispezione foglio principale
    res = inspect_document_content(file_bytes, "xlsx", "fatturato.xlsx")
    assert res["success"] is True
    assert res["active_sheet"] == "Fatturato 2026"
    assert "Clienti" in res["sheets"]
    assert res["total_rows"] == 4
    assert res["total_cols"] == 4
    assert "| Riga |" in res["markdown_table"]
    assert "A (Mese)" in res["markdown_table"]
    assert "Gennaio" in res["markdown_table"]

    # Ispezione con filtro per riga specifica (query)
    res_filtered = inspect_document_content(file_bytes, "xlsx", "fatturato.xlsx", query="Febbraio")
    assert res_filtered["success"] is True
    assert res_filtered["displayed_rows_count"] == 1
    assert "Febbraio" in res_filtered["markdown_table"]
    assert "Gennaio" not in res_filtered["markdown_table"]

    # Ispezione del secondo foglio
    res_sheet2 = inspect_document_content(file_bytes, "xlsx", "fatturato.xlsx", sheet_name="Clienti")
    assert res_sheet2["success"] is True
    assert res_sheet2["active_sheet"] == "Clienti"
    assert "Acme Corp" in res_sheet2["markdown_table"]


def test_inspect_document_content_csv():
    csv_data = "Data,Fornitore,Importo,Stato\n15/01/2026,Enel,85.50,Pagato\n20/02/2026,TIM,39.90,In Scadenza\n"
    res = inspect_document_content(csv_data.encode("utf-8"), "csv", "bollette.csv", query="TIM")
    assert res["success"] is True
    assert res["file_type"] == "csv"
    assert "TIM" in res["markdown_table"]
    assert "Enel" not in res["markdown_table"]


def test_inspect_document_content_docx():
    doc = docx.Document()
    doc.add_paragraph("Relazione Annuale di Bilancio")
    doc.add_paragraph("I costi di gestione sono diminuiti del 12% nel primo trimestre.")
    tbl = doc.add_table(rows=2, cols=2)
    tbl.rows[0].cells[0].text = "Voce"
    tbl.rows[0].cells[1].text = "Valore"
    tbl.rows[1].cells[0].text = "Spese IT"
    tbl.rows[1].cells[1].text = "4200 EUR"

    buf = io.BytesIO()
    doc.save(buf)
    file_bytes = buf.getvalue()

    res = inspect_document_content(file_bytes, "docx", "bilancio.docx", query="Spese IT")
    assert res["success"] is True
    assert "Spese IT" in res["content_text"]
    assert res["total_paragraphs"] >= 2
    assert res["total_tables"] == 1


def test_agent_service_execute_tool_read_vault_document_content(db_session: Session, tmp_path):
    # Salva un file Excel cifrato nel caveau
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spese Studio"
    ws.append(["Categoria", "Importo", "Note"])
    ws.append(["Cancelleria", 145.0, "Cartucce e carta"])
    ws.append(["Software", 590.0, "Licenze annuali"])

    buf = io.BytesIO()
    wb.save(buf)
    saved_path = save_uploaded_file(buf.getvalue(), "spese_studio.xlsx", target_dir=tmp_path)

    doc = Document(
        title="Spese Studio 2026",
        file_path=saved_path,
        file_type="xlsx",
        doc_type="generico",
        summary="Foglio Excel delle spese studio",
        thread_id="general"
    )
    db_session.add(doc)
    db_session.commit()

    agent = AgenticChatService()
    tool_res = agent.execute_tool(
        "read_vault_document_content",
        {"document_title": "Spese Studio 2026", "query": "Software"},
        db=db_session,
        thread_id="general"
    )

    assert tool_res["success"] is True
    assert tool_res["document_id"] == doc.id
    assert "Software" in tool_res["markdown_table"]
    assert "590" in tool_res["markdown_table"]


def test_agent_fallback_deterministic_excel_query(db_session: Session, tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fatture"
    ws.append(["Numero", "Cliente", "Totale"])
    ws.append(["FAT-01", "Mario Rossi", 1250.0])
    ws.append(["FAT-02", "Luigi Bianchi", 3400.0])

    buf = io.BytesIO()
    wb.save(buf)
    saved_path = save_uploaded_file(buf.getvalue(), "fatture_clienti.xlsx", target_dir=tmp_path)

    doc = Document(
        title="Fatture Clienti 2026",
        file_path=saved_path,
        file_type="xlsx",
        doc_type="generico",
        summary="Registro fatture clienti",
        thread_id="general"
    )
    db_session.add(doc)
    db_session.commit()

    agent = AgenticChatService()
    # Domanda specifica sui dati di Excel
    resp = agent.run_turn(
        user_text="cosa c'è nella riga del cliente Mario Rossi nel foglio excel?",
        db=db_session,
        thread_id="general"
    )

    assert resp.action == "read_vault_document_content"
    assert "Mario Rossi" in resp.reply
    assert "1250" in resp.reply


def test_mcp_server_tool_read_vault_document_content(db_session: Session, tmp_path, monkeypatch):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consuntivo"
    ws.append(["Voce", "Q1", "Q2"])
    ws.append(["Hardware", 3200, 4100])

    buf = io.BytesIO()
    wb.save(buf)
    saved_path = save_uploaded_file(buf.getvalue(), "consuntivo.xlsx", target_dir=tmp_path)

    doc = Document(
        title="Consuntivo Budget",
        file_path=saved_path,
        file_type="xlsx",
        doc_type="generico",
        summary="Budget consuntivo",
        thread_id="general"
    )
    db_session.add(doc)
    db_session.commit()

    # Monkeypatch _get_db_session in mcp_server per usare db_session di test
    from contextlib import contextmanager
    @contextmanager
    def mock_session():
        yield db_session

    monkeypatch.setattr(mcp_mod, "_get_db_session", mock_session)

    mcp_tool_res = mcp_mod.read_vault_document_content(document_title="Consuntivo Budget")
    assert "=== CONTENUTO DOCUMENTO: 'Consuntivo Budget' ===" in mcp_tool_res
    assert "Hardware" in mcp_tool_res
    assert "3200" in mcp_tool_res
