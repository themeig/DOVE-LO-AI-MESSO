"""Test per gli strumenti Zip, Unzip e lettura file Word (.docx) ed Excel (.xlsx, .csv)."""
import io
import json
import zipfile
import pytest
import docx
import openpyxl
from fastapi.testclient import TestClient

from app.main import app
from app.models.database import init_db, Document, get_session_maker, get_engine
from app.services.ai_service import get_ai_service, MockAIService
from app.services.archive_service import (
    create_zip_from_documents,
    unzip_document_to_vault,
    extract_text_from_office_file,
    render_office_file_to_html
)
from app.services.agent_service import AgenticChatService
from app.services.crypto_service import get_vault_manager

client = TestClient(app)

def setup_module():
    init_db()
    mgr = get_vault_manager()
    mgr.initialize_if_needed("1234")


def create_dummy_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("Contratto di Consulenza Software", 0)
    doc.add_paragraph("Il presente contratto e stipulato tra Azienda Alfa e Professionista Mario Rossi.")
    doc.add_paragraph("Compenso pattuito: 3.500,00 Euro con scadenza pagamento 2026-12-31.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def create_dummy_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Spese e Scadenze"
    ws.append(["Descrizione", "Fornitore", "Importo Euro", "Data Scadenza"])
    ws.append(["Bolletta Fibra Ottica", "TIM Business", 120.50, "2026-11-15"])
    ws.append(["Assicurazione Ufficio", "Generali Assicurazioni", 450.00, "2026-12-01"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def create_dummy_csv() -> bytes:
    content = "Fornitore,Tipo,Importo,Scadenza\nA2A Energia,Bolletta Luce,85.20,2026-10-30\nAcquedotto,Acqua,42.00,2026-11-05\n"
    return content.encode("utf-8")


def create_dummy_zip() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("fattura_fornitore.txt", "Fattura numero 104 per consulenza software importo 1200 euro scadenza 2026-11-30")
        zf.writestr("riassunto_spese.csv", "Descrizione,Importo\nInternet,50\nCancelleria,25\n")
    return buf.getvalue()


def test_office_extraction_docx():
    """Test estrazione testo da documento Word (.docx)."""
    docx_bytes = create_dummy_docx()
    text = extract_text_from_office_file(docx_bytes, "contratto.docx")
    assert "Contratto di Consulenza" in text
    assert "Mario Rossi" in text
    assert "3.500" in text

    # Estrazione AI da file docx
    ai = MockAIService()
    ext = ai.extract_document(docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename="contratto_consulenza.docx")
    assert ext.doc_type in ["contratto", "documento_word", "generico"]
    assert ext.title


def test_office_extraction_xlsx():
    """Test estrazione testo/tabelle da foglio Excel (.xlsx)."""
    xlsx_bytes = create_dummy_xlsx()
    text = extract_text_from_office_file(xlsx_bytes, "spese_2026.xlsx")
    assert "TIM Business" in text
    assert "120.5" in text
    assert "Generali Assicurazioni" in text

    # Estrazione AI da file xlsx
    ai = MockAIService()
    ext = ai.extract_document(xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", filename="spese_aziendali.xlsx")
    assert ext.doc_type in ["foglio_calcolo", "spese", "fattura", "generico"]
    assert ext.title


def test_office_extraction_csv():
    """Test estrazione dati da file CSV."""
    csv_bytes = create_dummy_csv()
    text = extract_text_from_office_file(csv_bytes, "bollette.csv")
    assert "A2A Energia" in text
    assert "85.20" in text


def test_create_and_unzip_archive_service():
    """Test creazione archivio ZIP da documenti del caveau e successivo unzip con re-indicizzazione."""
    from app.models.database import get_db
    get_db_func = app.dependency_overrides.get(get_db, get_db)
    db = next(get_db_func())
    try:
        # 1. Carica 2 documenti di prova
        pdf_bytes = b"%PDF-1.4 documento da zippare"
        res_up1 = client.post(
            "/api/documents/upload",
            files={"file": ("doc_uno.pdf", io.BytesIO(pdf_bytes), "application/pdf")},
            data={"thread_id": "general"}
        )
        assert res_up1.status_code == 201
        doc1_id = res_up1.json()["document_id"]

        docx_bytes = create_dummy_docx()
        res_up2 = client.post(
            "/api/documents/upload",
            files={"file": ("contratto_alfa.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            data={"thread_id": "general"}
        )
        assert res_up2.status_code == 201
        doc2_id = res_up2.json()["document_id"]

        # 2. Crea archivio ZIP con i 2 documenti
        zip_doc = create_zip_from_documents(
            db=db,
            document_ids=[doc1_id, doc2_id],
            archive_title="Archivio Contratti e PDF",
            thread_id="general"
        )
        assert zip_doc is not None
        assert zip_doc.id > 0
        assert zip_doc.file_type == "zip"
        assert zip_doc.doc_type == "archivio_zip"

        # 3. Esegui unzip dell'archivio appena creato
        extracted_docs = unzip_document_to_vault(
            db=db,
            document_id=zip_doc.id,
            thread_id="general"
        )
        assert len(extracted_docs) >= 2
    finally:
        db.close()


def test_agent_zip_and_unzip_tools():
    """Test esecuzione dei tool create_zip_archive e unzip_vault_archive dell'agente."""
    from app.models.database import get_db
    get_db_func = app.dependency_overrides.get(get_db, get_db)
    db = next(get_db_func())
    try:
        agent = AgenticChatService()

        # Upload di un file per il test
        res_up = client.post(
            "/api/documents/upload",
            files={"file": ("tessera_test.pdf", io.BytesIO(b"%PDF-1.4 tessera sanitaria test"), "application/pdf")},
            data={"thread_id": "general"}
        )
        assert res_up.status_code == 201
        doc_id = res_up.json()["document_id"]

        # 1. Tool create_zip_archive
        zip_res = agent.execute_tool(
            "create_zip_archive",
            {"query": "tessera", "archive_name": "Pacchetto_Tessere.zip"},
            db=db,
            thread_id="general"
        )
        assert zip_res.get("success") is True
        assert zip_res.get("zip_document_id") is not None
        zip_id = zip_res["zip_document_id"]

        # 2. Tool unzip_vault_archive
        unzip_res = agent.execute_tool(
            "unzip_vault_archive",
            {"document_id": zip_id},
            db=db,
            thread_id="general"
        )
        assert unzip_res.get("success") is True
        assert len(unzip_res.get("extracted_documents", [])) >= 1
    finally:
        db.close()


def test_render_office_file_to_html_docx():
    """Verifica la generazione di HTML per documenti Word (.docx)."""
    docx_bytes = create_dummy_docx()
    result = render_office_file_to_html(docx_bytes, "contratto_consulenza.docx")
    assert result["success"] is True
    assert result["format"] == "word"
    assert "Contratto di Consulenza Software" in result["html_content"]
    assert "Mario Rossi" in result["html_content"]
    assert "<h1" in result["html_content"] or "<p" in result["html_content"]


def test_render_office_file_to_html_xlsx_multisheet():
    """Verifica il rendering HTML con schede multiple per fogli Excel (.xlsx)."""
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Fatture Entrata"
    ws1.append(["Cliente", "Importo", "Stato"])
    ws1.append(["Acme SpA", 1500.0, "Incassato"])

    ws2 = wb.create_sheet(title="Fatture Uscita")
    ws2.append(["Fornitore", "Importo", "Scadenza"])
    ws2.append(["Enel Energia", 230.40, "2026-11-20"])

    buf = io.BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    result = render_office_file_to_html(xlsx_bytes, "bilancio_trimestrale.xlsx")
    assert result["success"] is True
    assert result["format"] == "excel"
    assert len(result["sheets"]) == 2
    assert result["sheets"][0]["name"] == "Fatture Entrata"
    assert result["sheets"][1]["name"] == "Fatture Uscita"
    assert "Acme SpA" in result["html_content"]
    assert "Enel Energia" in result["html_content"]
    # Verifica presenza dei pulsanti tab per il cambio scheda
    assert "switchOfficeSheet(0)" in result["html_content"]
    assert "switchOfficeSheet(1)" in result["html_content"]


def test_render_office_file_to_html_csv():
    """Verifica il rendering HTML per file CSV."""
    csv_bytes = create_dummy_csv()
    result = render_office_file_to_html(csv_bytes, "scadenze.csv")
    assert result["success"] is True
    assert result["format"] == "excel"
    assert len(result["sheets"]) == 1
    assert "A2A Energia" in result["html_content"]
    assert "85.20" in result["html_content"]


def test_document_preview_content_endpoint():
    """Verifica l'endpoint GET /api/documents/preview-content."""
    # 1. Upload di un file Excel
    xlsx_bytes = create_dummy_xlsx()
    res_up_xls = client.post(
        "/api/documents/upload",
        files={"file": ("report_costi.xlsx", io.BytesIO(xlsx_bytes), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
        data={"thread_id": "general"}
    )
    assert res_up_xls.status_code == 201
    xls_doc_id = res_up_xls.json()["document_id"]
    xls_file_url = res_up_xls.json()["file_url"]

    # 2. Richiesta anteprima tramite document_id
    res_prev1 = client.get(f"/api/documents/preview-content?document_id={xls_doc_id}")
    assert res_prev1.status_code == 200
    data1 = res_prev1.json()
    assert data1["success"] is True
    assert data1["format"] == "excel"
    assert "TIM Business" in data1["html_content"]

    # 3. Richiesta anteprima tramite file_url
    res_prev2 = client.get(f"/api/documents/preview-content?file_url={xls_file_url}")
    assert res_prev2.status_code == 200
    data2 = res_prev2.json()
    assert data2["success"] is True
    assert data2["format"] == "excel"

    # 4. Upload ed anteprima di un file Word (.docx)
    docx_bytes = create_dummy_docx()
    res_up_doc = client.post(
        "/api/documents/upload",
        files={"file": ("accordo_quadro.docx", io.BytesIO(docx_bytes), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        data={"thread_id": "general"}
    )
    assert res_up_doc.status_code == 201
    doc_id = res_up_doc.json()["document_id"]

    res_prev3 = client.get(f"/api/documents/preview-content?document_id={doc_id}")
    assert res_prev3.status_code == 200
    data3 = res_prev3.json()
    assert data3["success"] is True
    assert data3["format"] == "word"
    assert "Contratto di Consulenza Software" in data3["html_content"]

    # 5. Documento inesistente
    res_404 = client.get("/api/documents/preview-content?document_id=999999")
    assert res_404.status_code == 404


def test_office_extraction_and_rendering_xls_legacy():
    """Verifica l'estrazione e rendering di file Excel legacy (.xls) tramite xlrd."""
    from unittest.mock import MagicMock, patch

    mock_sheet = MagicMock()
    mock_sheet.name = "Spese Legacy"
    mock_sheet.nrows = 3
    mock_sheet.ncols = 3
    mock_sheet.cell_value.side_effect = lambda r, c: [
        ["Descrizione", "Importo", "Data"],
        ["Affitto Ufficio", 850.0, "2026-11-01"],
        ["Cancelleria", 45.0, "2026-11-05"]
    ][r][c]

    mock_wb = MagicMock()
    mock_wb.sheets.return_value = [mock_sheet]
    mock_wb.sheet_by_index.return_value = mock_sheet

    with patch("xlrd.open_workbook", return_value=mock_wb):
        text = extract_text_from_office_file(b"dummy_xls_content", "spese.xls")
        assert "Spese Legacy" in text
        assert "Affitto Ufficio" in text
        assert "850" in text

        rendered = render_office_file_to_html(b"dummy_xls_content", "spese.xls")
        assert rendered["success"] is True
        assert rendered["format"] == "excel"
        assert "Affitto Ufficio" in rendered["html_content"]
        assert "Spese Legacy" in rendered["sheets"][0]["name"]


def test_render_office_file_to_html_zip():
    """Verifica il rendering dell'anteprima esplorabile per archivi ZIP."""
    zip_bytes = create_dummy_zip()
    res = render_office_file_to_html(zip_bytes, "documenti_progetto.zip")
    assert res["success"] is True
    assert res["format"] == "zip"
    assert res["file_count"] == 2
    assert "fattura_fornitore.txt" in res["html_content"]
    assert "riassunto_spese.csv" in res["html_content"]
    assert "unzipCurrentModalArchive()" in res["html_content"]
    assert "Estrai tutti i file nel Caveau" in res["html_content"]


def test_unzip_document_endpoint_post():
    """Verifica gli endpoint REST POST /api/documents/{document_id}/unzip e POST /api/documents/unzip."""
    zip_bytes = create_dummy_zip()
    res_up = client.post(
        "/api/documents/upload",
        files={"file": ("archivio_da_scompattare.zip", io.BytesIO(zip_bytes), "application/zip")},
        data={"thread_id": "general"}
    )
    assert res_up.status_code == 201
    zip_id = res_up.json()["document_id"]

    # 1. Unzip via POST /{document_id}/unzip
    res_unzip = client.post(f"/api/documents/{zip_id}/unzip")
    assert res_unzip.status_code == 200
    data = res_unzip.json()
    assert data["success"] is True
    assert data["count"] == 2
    assert len(data["documents"]) == 2
    doc_titles = [d["title"] for d in data["documents"]]
    assert any("fattura" in t.lower() for t in doc_titles)
    assert any("spese" in t.lower() or "riassunto" in t.lower() for t in doc_titles)

    # 2. Verifica che i documenti estratti abbiano un download url valido
    first_id = data["documents"][0]["id"]
    res_dl = client.get(f"/api/documents/{first_id}/download")
    assert res_dl.status_code == 200
    assert len(res_dl.content) > 0


def test_chat_unzip_turn_interception():
    """Verifica che l'assistente riconosca i comandi conversazionali di scompattamento ZIP."""
    from app.models.database import get_db
    get_db_func = app.dependency_overrides.get(get_db, get_db)
    db = next(get_db_func())
    try:
        zip_bytes = create_dummy_zip()
        res_up = client.post(
            "/api/documents/upload",
            files={"file": ("archivio_chat_test.zip", io.BytesIO(zip_bytes), "application/zip")},
            data={"thread_id": "chat_unzip_test"}
        )
        assert res_up.status_code == 201

        agent = AgenticChatService()
        turn_res = agent.run_turn(
            user_text="scompatta lo zip per favore",
            db=db,
            thread_id="chat_unzip_test"
        )
        assert "scompattato" in turn_res.reply.lower() or "estratti" in turn_res.reply.lower()
        assert turn_res.documents is not None
        assert len(turn_res.documents) >= 2
        # Verifica che i documenti siano presentati uno per uno con numerazione e dettagli
        assert "1." in turn_res.reply
        assert any(d.get("category_label") for d in turn_res.documents)
    finally:
        db.close()


def test_unzip_calls_ai_service_for_each_file(monkeypatch):
    """Verifica che la decompressione ZIP analizzi ciascun file estratto tramite il motore AI."""
    from unittest.mock import MagicMock
    from app.models.database import get_db
    from app.models.schemas import ExtractedDocument

    mock_ai = MagicMock()
    mock_ai.extract_document.side_effect = lambda file_bytes, mime_type, filename: ExtractedDocument(
        title=f"AI Title {filename}",
        doc_type="bolletta" if "fattura" in filename else "contratto",
        issuer="AI Test Issuer",
        amount=99.90 if "fattura" in filename else None,
        due_date="2026-11-20" if "fattura" in filename else None,
        summary=f"Analisi AI dettagliata per {filename}",
        category="utenze_bollette" if "fattura" in filename else "contratti_polizze",
        category_label="Utenze & Bollette" if "fattura" in filename else "Contratti, Polizze & Assicurazioni",
        category_icon="fa-bolt"
    )

    import app.services.archive_service
    monkeypatch.setattr(app.services.archive_service, "get_ai_service", lambda db=None: mock_ai)

    engine = get_engine()
    SessionLocal = get_session_maker(engine)
    with SessionLocal() as db:
        zip_bytes = create_dummy_zip()
        from app.services.document_service import save_uploaded_file
        saved_zip_path = save_uploaded_file(zip_bytes, "test_mock_ai.zip")
        zip_doc = Document(
            title="Archivio Test Mock AI.zip",
            file_path=saved_zip_path,
            file_type="zip",
            doc_type="archivio_zip",
            summary="Test zip",
            thread_id="test_ai_unzip"
        )
        db.add(zip_doc)
        db.commit()
        db.refresh(zip_doc)

        extracted = unzip_document_to_vault(db=db, document_id=zip_doc.id, thread_id="test_ai_unzip")
        assert len(extracted) == 2
        # L'AI deve essere stata invocata per ciascun file
        assert mock_ai.extract_document.call_count == 2
        assert all("AI Title" in d.title for d in extracted)
        assert any(d.amount == 99.90 for d in extracted)
        assert any(d.category_label == "Utenze & Bollette" for d in extracted)



